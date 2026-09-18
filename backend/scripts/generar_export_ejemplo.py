"""Genera `fixtures/export_ejemplo.xlsx`, un export sintético acorde a `mapeo_columnas.json`.

Incluye: ítems con código del catálogo semilla, un código inexistente (XXX-999), un ítem sin
código y una fila de total que el parser debe omitir.

Uso:  python scripts/generar_export_ejemplo.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

RAIZ = Path(__file__).resolve().parent.parent
DESTINO = RAIZ / "fixtures" / "export_ejemplo.xlsx"

FILAS = [
    # código, descripción, cantidad, precio unitario
    ("SIL-001", "Silla Tiffany blanca", 120, 45.00),
    ("SIL-004", "Silla Crossback natural", 80, 65.00),
    ("MES-002", "Mesa redonda 1.80 m", 12, 250.00),
    ("MES-005", "Mesa imperial madera 3 m", 4, 900.00),
    ("LOU-001", "Sala lounge lino crudo (3 piezas)", 3, 1800.00),
    ("BAR-002", "Barra de bebidas madera 2.4 m", 1, 3500.00),
    ("ILU-003", "Lámpara de pie Edison", 6, 320.00),
    ("MAN-001", "Mantel lino arena 3 m", 12, 120.00),
    ("XXX-999", "Pérgola con cortinas de gasa", 1, 4500.00),  # código inexistente
    ("", "Letrero personalizado con nombre de los novios", 1, 2200.00),  # sin código
    ("DEC-004", "Centro de mesa floral bajo", 12, 450.00),
    ("TAR-001", "Tarima 6 x 4 m con faldón", 1, 5200.00),
]


def generar(destino: Path = DESTINO) -> Path:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Cotización"

    hoja["A1"] = "COTIZACIÓN DE MOBILIARIO — Minimal 4.0"
    hoja["A1"].font = Font(bold=True, size=13)
    hoja["A2"] = "Cliente:"
    hoja["B2"] = "Hacienda San Pedro Eventos"
    hoja["A3"] = "Referencia:"
    hoja["B3"] = "COT-2026-0142"

    encabezados = ["Código", "Descripción", "Cantidad", "Precio", "Importe"]
    for columna, texto in enumerate(encabezados, start=1):
        celda = hoja.cell(row=5, column=columna, value=texto)
        celda.font = Font(bold=True)

    fila_actual = 6
    for codigo, descripcion, cantidad, precio in FILAS:
        hoja.cell(row=fila_actual, column=1, value=codigo or None)
        hoja.cell(row=fila_actual, column=2, value=descripcion)
        hoja.cell(row=fila_actual, column=3, value=cantidad)
        hoja.cell(row=fila_actual, column=4, value=precio)
        hoja.cell(row=fila_actual, column=5, value=round(cantidad * precio, 2))
        fila_actual += 1

    hoja.cell(row=fila_actual + 1, column=2, value="Total")
    hoja.cell(row=fila_actual + 1, column=5, value=round(sum(c * p for _, _, c, p in FILAS), 2))

    hoja.column_dimensions["A"].width = 12
    hoja.column_dimensions["B"].width = 48
    for letra in "CDE":
        hoja.column_dimensions[letra].width = 12

    destino.parent.mkdir(parents=True, exist_ok=True)
    libro.save(destino)
    return destino


if __name__ == "__main__":
    print(f"Generado: {generar()}")
